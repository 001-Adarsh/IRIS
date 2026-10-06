from pathlib import Path
import re

from docx import Document


DATA_DIR = Path(__file__).resolve().parent.parent / "data"


def _filename_only(filename: str) -> str:
    name = Path(filename or "iris_document.docx").name
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(name).stem).strip(" .")
    return f"{stem or 'iris_document'}.docx"


def create_document(
    filename: str = "iris_document.docx",
    title: str = "IRIS Document",
    content: str = "",
) -> str:
    """Create a Word document in IRIS's data directory and return its path."""
    if not isinstance(content, str) or not content.strip():
        raise ValueError("Document content must be a non-empty string.")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(title.strip() or "IRIS Document", level=0)

    for line in content.splitlines():
        text = line.strip()
        if not text:
            continue
        if text.startswith("### "):
            document.add_heading(text[4:], level=3)
        elif text.startswith("## "):
            document.add_heading(text[3:], level=2)
        elif text.startswith("# "):
            document.add_heading(text[2:], level=1)
        elif re.match(r"^[-*]\s+", text):
            document.add_paragraph(re.sub(r"^[-*]\s+", "", text), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", text):
            document.add_paragraph(
                re.sub(r"^\d+[.)]\s+", "", text), style="List Number"
            )
        else:
            document.add_paragraph(text)

    path = DATA_DIR / _filename_only(filename)
    document.save(path)
    return str(path)
