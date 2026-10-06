import os
from docx import Document

class DocxCompiler:
    """
    Compiler tool for generating professional .docx documents from structured content or strings.
    Integrated into IRIS Autonomous Core to fulfill documentation generation requests.
    """
    def __init__(self):
        pass

    def create_document(self, title: str, sections: list, output_path: str = "output.docx") -> str:
        """
        Creates a Word document with a given title and a list of section dicts (heading, body).
        Returns the absolute file path of the generated document.
        """
        doc = Document()
        doc.add_heading(title, level=0)

        for section in sections:
            heading = section.get("heading")
            body = section.get("body")
            
            if heading:
                doc.add_heading(heading, level=1)
            if body:
                doc.add_paragraph(body)

        abs_path = os.path.abspath(output_path)
        doc.save(abs_path)
        return abs_path

    def compile_from_text(self, title: str, content: str, output_path: str = "output.docx") -> str:
        """
        Takes raw text content, breaks it down into paragraphs, and generates a structured .docx file.
        """
        doc = Document()
        doc.add_heading(title, level=0)
        
        paragraphs = content.split("\n\n")
        for p in paragraphs:
            if p.strip():
                doc.add_paragraph(p.strip())

        abs_path = os.path.abspath(output_path)
        doc.save(abs_path)
        return abs_path