from pathlib import Path
import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


DATA_DIR = Path(__file__).resolve().parent.parent / "data"


def read_txt(file_path):
    return Path(file_path).read_text(
        encoding="utf-8",
        errors="replace"
    )


def read_csv(file_path):
    df = pd.read_csv(file_path)
    return df.to_string(index=False)


def read_xlsx(file_path):
    sheets = pd.read_excel(
        file_path,
        sheet_name=None
    )

    output = []

    for sheet_name, df in sheets.items():
        output.append(
            f"=== SHEET: {sheet_name} ==="
        )
        output.append(
            df.to_string(index=False)
        )

    return "\n\n".join(output)


def read_pdf(file_path):
    reader = PdfReader(file_path)

    pages = []

    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""

        pages.append(
            f"=== PAGE {i} ===\n{text}"
        )

    return "\n\n".join(pages)


def read_docx(file_path):
    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(
                paragraph.text.strip()
            )

    return "\n".join(paragraphs)


def read_pptx(file_path):
    presentation = Presentation(file_path)

    slides = []

    for i, slide in enumerate(
        presentation.slides,
        1
    ):

        slide_text = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_text.append(text)

        slides.append(
            f"=== SLIDE {i} ===\n"
            + "\n".join(slide_text)
        )

    return "\n\n".join(slides)


def read_document(filename):
    file_path = Path(filename)

    if not file_path.is_absolute():
        file_path = DATA_DIR / filename

    if not file_path.exists():
        return f"File not found: {file_path}"

    extension = file_path.suffix.lower()

    try:

        if extension == ".txt":
            return read_txt(file_path)

        elif extension == ".csv":
            return read_csv(file_path)

        elif extension in [".xlsx", ".xls"]:
            return read_xlsx(file_path)

        elif extension == ".pdf":
            return read_pdf(file_path)

        elif extension == ".docx":
            return read_docx(file_path)

        elif extension == ".pptx":
            return read_pptx(file_path)

        else:
            return (
                f"Unsupported file type: "
                f"{extension}"
            )

    except Exception as e:

        return (
            f"Document reading error: {e}"
        )