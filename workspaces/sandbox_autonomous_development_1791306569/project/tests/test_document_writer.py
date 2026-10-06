import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from core.agent import INDIAN_STATES, IRISAgent
from core.executor import IRISExecutor
from core.task_verifier import IRISTaskVerifier
from tools.manager import ToolManager
from tools.document_writer import create_document


class TestDocumentWriter(unittest.TestCase):
    def test_create_word_document(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("tools.document_writer.DATA_DIR", Path(directory)):
                path = Path(
                    create_document(
                        filename="facts.txt",
                        title="Sample facts",
                        content="# Details\n1. First fact\n- Second fact",
                    )
                )

            self.assertEqual(path.suffix, ".docx")
            self.assertTrue(path.is_file())
            text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
            self.assertIn("Sample facts", text)
            self.assertIn("Details", text)
            self.assertIn("First fact", text)
            self.assertIn("Second fact", text)

    def test_registered_document_tool_executes_and_verifies(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("tools.document_writer.DATA_DIR", Path(directory)):
                manager = ToolManager()
                arguments = {
                    "filename": "verified.docx",
                    "title": "Verified",
                    "content": "A saved document.",
                }
                execution = IRISExecutor(manager).execute_plan(
                    {
                        "tools": [
                            {
                                "name": "create_document",
                                "arguments": arguments,
                            }
                        ]
                    }
                )

            self.assertTrue(execution["success"])
            tool_result = execution["results"][0]["result"]
            verification = IRISTaskVerifier().verify(
                "create a Word document",
                "create_document",
                arguments,
                tool_result,
            )
            self.assertTrue(verification["passed"])
            self.assertTrue(Path(tool_result).is_file())

    def test_agent_creates_document_for_indian_states_without_web_search(self):
        with tempfile.TemporaryDirectory() as directory:
            class StubRouter:
                providers = {}

                def compare(self, prompt, timeout=8.0):
                    raise AssertionError(
                        f"Unexpected provider call: {prompt} (timeout={timeout})."
                    )

            class StubTools:
                def execute(self, tool_name, **kwargs):
                    self.tool_name = tool_name
                    return create_document(**kwargs)

            agent = object.__new__(IRISAgent)
            agent.project_root = Path(directory)
            agent.chat_history = []
            agent.last_objective = ""
            agent.last_error_log = ""
            agent.router = StubRouter()
            agent.tools = StubTools()

            with patch("tools.document_writer.DATA_DIR", Path(directory) / "data"):
                reply = agent.run(
                    "give me a doc which contain all state of india with it's name"
                )

            self.assertIn("Word document created:", reply)
            self.assertEqual(agent.tools.tool_name, "create_document")
            output_path = Path(reply.split(": ", 1)[1])
            self.assertTrue(output_path.is_file())
            document_text = "\n".join(
                paragraph.text for paragraph in Document(output_path).paragraphs
            )
            for state in ("Andhra Pradesh", "West Bengal", "Uttarakhand"):
                self.assertIn(state, document_text)
            self.assertNotIn("Jammu and Kashmir", document_text)
            self.assertEqual(
                sum(state in document_text for state in INDIAN_STATES), 28
            )


if __name__ == "__main__":
    unittest.main()
